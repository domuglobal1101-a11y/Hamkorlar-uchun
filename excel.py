"""Excel (.xlsx) hisobot: Hamkorlar, Savdolar, To'lovlar, Sotuvchilar varaqlari."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import db

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
MONEY = '#,##0'


def _sheet(wb, title, headers, rows, money_cols=()):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for r in rows:
        ws.append(list(r))
    for col in money_cols:
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].number_format = MONEY
    for i, h in enumerate(headers, 1):
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:200]]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 40)
    ws.freeze_panes = "A2"
    return ws


def build(date_from=None, date_to=None):
    """Excel faylni bayt ko'rinishida qaytaradi. Savdo/to'lovlar sana bo'yicha filtrlanadi."""
    wb = Workbook()
    wb.remove(wb.active)

    clients = db.clients_with_balance()
    _sheet(
        wb, "Hamkorlar",
        ["ID", "Ism", "Telefon", "Yo'nalish", "Jami aylanma", "Bosqich %", "Jami bonus",
         "To'langan", "Qoldiq", "Savdolar", "Telegram", "Ro'yxatdan o'tgan"],
        [
            (c["id"], c["full_name"], c["phone"], c["category"], c["total_turnover"],
             db.percent_for(c["total_turnover"]), c["total_bonus"], c["paid"],
             c["total_bonus"] - c["paid"], c["sales_count"], "ha" if c["telegram_id"] else "yo'q",
             (c["created_at"] or "").replace("T", " "))
            for c in clients
        ],
        money_cols=(5, 7, 8, 9),
    )

    sales = db.list_sales(date_from, date_to)
    _sheet(
        wb, "Savdolar",
        ["ID", "Sana", "Hamkor", "Telefon", "Summa", "Foiz", "Bonus", "Aylanma (keyin)",
         "Xaridor", "Sotuvchi", "Holat"],
        [
            (s["id"], (s["created_at"] or "").replace("T", " "), s["client_name"], s["client_phone"],
             s["amount"], s["percent"], s["bonus"], s["turnover_after"], s["customer"],
             s["seller_name"], "faol" if s["status"] == "active" else "bekor qilingan")
            for s in sales
        ],
        money_cols=(5, 7, 8),
    )

    payouts = db.list_payouts(date_from=date_from, date_to=date_to)
    _sheet(
        wb, "To'lovlar",
        ["ID", "Sana", "Hamkor", "Telefon", "Summa", "Izoh", "Kim kiritdi"],
        [
            (p["id"], (p["created_at"] or "").replace("T", " "), p["client_name"], p["client_phone"],
             p["amount"], p["note"], p["created_by"])
            for p in payouts
        ],
        money_cols=(5,),
    )

    sellers = db.sellers_stats(date_from, date_to)
    _sheet(
        wb, "Sotuvchilar",
        ["ID", "Ism", "Savdolar", "Aylanma", "Bonus"],
        [(s["id"], s["name"], s["cnt"], s["turn"], s["bon"]) for s in sellers],
        money_cols=(4, 5),
    )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
